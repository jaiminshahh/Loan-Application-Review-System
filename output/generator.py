"""
LAYER 5: Output Generator

Pure Python - creates Excel decision table and JSON report.
"""
import json
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.data_extractor import ExtractedData
from core.evaluators import Evaluation
from core.aggregator import AggregatedResults
from core.decision_engine import FinalDecision
from config.factors import Status


class OutputGenerator:
    """Generates output files from evaluation results."""
    
    def __init__(self, output_dir: str = "output"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)
        
        # Styles
        self.title_font = Font(bold=True, size=16, color="FFFFFF")
        self.title_fill = PatternFill("solid", fgColor="1F4E79")
        self.header_font = Font(bold=True, size=10, color="FFFFFF")
        self.header_fill = PatternFill("solid", fgColor="2F5496")
        self.section_font = Font(bold=True, size=11, color="1F4E79")
        self.section_fill = PatternFill("solid", fgColor="D6DCE4")
        self.pass_fill = PatternFill("solid", fgColor="C6EFCE")
        self.pass_font = Font(bold=True, color="006100")
        self.fail_fill = PatternFill("solid", fgColor="FFC7CE")
        self.fail_font = Font(bold=True, color="9C0006")
        self.review_fill = PatternFill("solid", fgColor="FFEB9C")
        self.review_font = Font(bold=True, color="9C5700")
        self.border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
    
    def generate_excel(
        self,
        data: ExtractedData,
        results: AggregatedResults,
        decision: FinalDecision
    ) -> str:
        """Generate Excel decision table."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Decision Table"
        
        row = 1
        
        # Title
        ws.merge_cells('A1:H1')
        ws['A1'] = "LOAN APPLICATION DECISION TABLE"
        ws['A1'].font = self.title_font
        ws['A1'].fill = self.title_fill
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30
        
        # Applicant info
        row = 2
        ws.merge_cells('A2:H2')
        info = f"Applicant: {data.applicant_name}  |  Loan: ${data.loan_amount:,.0f} ({data.loan_term_months} mo)  |  Purpose: {data.loan_purpose}"
        ws['A2'] = info
        ws['A2'].font = Font(bold=True, size=10)
        ws['A2'].fill = PatternFill("solid", fgColor="F2F2F2")
        ws['A2'].alignment = Alignment(horizontal='center')
        
        # Headers
        row = 4
        headers = ["Factor", "Value", "Status", "Risk", "Conf", "Source", "Weight", "Notes"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.border
        
        # Evaluations by category
        row = 5
        current_category = None
        
        for eval in results.evaluations:
            # Category header
            if eval.category != current_category:
                current_category = eval.category
                ws.merge_cells(f'A{row}:H{row}')
                cell = ws.cell(row=row, column=1, value=current_category.upper())
                cell.font = self.section_font
                cell.fill = self.section_fill
                cell.border = self.border
                row += 1
            
            # Factor row (now includes risk_score, confidence, source)
            source_display = eval.source
            if eval.disagreement:
                source_display += " (!)"
            
            values = [
                eval.factor_name,
                eval.value,
                eval.status.value,
                f"{eval.risk_score:.1f}",
                f"{eval.confidence:.2f}",
                source_display,
                f"{eval.weight * 100:.0f}%",
                eval.notes  # Full notes without truncation
            ]
            
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = self.border
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                
                # Style status column
                if col == 3:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    if value == "PASS":
                        cell.fill = self.pass_fill
                        cell.font = self.pass_font
                    elif value == "FAIL":
                        cell.fill = self.fail_fill
                        cell.font = self.fail_font
                    else:
                        cell.fill = self.review_fill
                        cell.font = self.review_font
                elif col == 5:  # Confidence column
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif col == 6:  # Source column
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    if eval.disagreement:
                        cell.font = Font(bold=True, color="FF0000")
            
            # Auto-adjust row height based on notes length (minimum 35, increase for longer notes)
            notes_lines = len(eval.notes) / 60  # Approximate lines needed (60 chars per line)
            ws.row_dimensions[row].height = max(35, min(150, 35 + notes_lines * 15))
            row += 1
        
        # Summary section
        row += 1
        ws.merge_cells(f'A{row}:H{row}')
        ws.cell(row=row, column=1, value="EVALUATION SUMMARY").font = self.section_font
        ws.cell(row=row, column=1).fill = self.section_fill
        ws.cell(row=row, column=1).border = self.border
        
        row += 1
        summary = [
            f"Total: {results.total_factors}",
            f"PASS: {results.pass_count}",
            f"REVIEW: {results.review_count}",
            f"FAIL: {results.fail_count}",
            f"Score: {results.weighted_score}%",
            f"Risk: {results.average_risk_score:.1f}",
            f"Conf: {results.average_confidence:.2f}",
            f"Level: {decision.risk_level}"
        ]
        for col, value in enumerate(summary, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = self.border
            cell.font = Font(bold=True)
            if "PASS" in value:
                cell.fill = self.pass_fill
            elif "FAIL" in value:
                cell.fill = self.fail_fill
            elif "REVIEW" in value:
                cell.fill = self.review_fill
        
        # Decision section
        row += 2
        ws.merge_cells(f'A{row}:H{row}')
        ws.cell(row=row, column=1, value="FINAL DECISION").font = self.section_font
        ws.cell(row=row, column=1).fill = self.section_fill
        ws.cell(row=row, column=1).border = self.border
        
        row += 1
        ws.merge_cells(f'A{row}:B{row}')
        cell = ws.cell(row=row, column=1, value=decision.decision)
        cell.font = Font(bold=True, size=14)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = self.border
        ws.row_dimensions[row].height = 35
        
        if decision.decision == "APPROVED":
            cell.font = Font(bold=True, size=14, color="006100")
            cell.fill = self.pass_fill
        elif decision.decision == "REJECTED":
            cell.font = Font(bold=True, size=14, color="9C0006")
            cell.fill = self.fail_fill
        else:
            cell.font = Font(bold=True, size=14, color="9C5700")
            cell.fill = self.review_fill
        
        # Rule applied
        ws.merge_cells(f'C{row}:H{row}')
        ws.cell(row=row, column=3, value=f"Rule: {decision.rule_applied} - {decision.rule_description}")
        ws.cell(row=row, column=3).border = self.border
        
        # Recommendation
        row += 1
        ws.merge_cells(f'A{row}:H{row}')
        cell = ws.cell(row=row, column=1, value=f"RECOMMENDATION: {decision.recommendation}")
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = self.border
        ws.row_dimensions[row].height = 50
        
        # Column widths (increased Notes column from 45 to 60 for better visibility)
        widths = [20, 20, 10, 8, 8, 12, 8, 60]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        # Save
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"decision_table_{timestamp}.xlsx"
        filepath = self.output_dir / filename
        wb.save(filepath)
        
        return str(filepath)
    
    def generate_json(
        self,
        data: ExtractedData,
        results: AggregatedResults,
        decision: FinalDecision
    ) -> str:
        """Generate JSON report."""
        report = {
            "generated_at": datetime.now().isoformat(),
            "applicant": {
                "name": data.applicant_name,
                "loan_amount": data.loan_amount,
                "loan_term_months": data.loan_term_months,
                "loan_purpose": data.loan_purpose
            },
            "extracted_data": {
                "credit_score": data.credit_score,
                "payment_history_pct": data.payment_history_pct,
                "credit_utilization_pct": data.credit_utilization_pct,
                "hard_inquiries": data.hard_inquiries,
                "employment_status": data.employment_status,
                "employment_duration_years": data.employment_duration_years,
                "annual_income": data.annual_income,
                "monthly_income": data.monthly_income,
                "monthly_debt_payments": data.monthly_debt_payments,
                "dti_ratio": data.dti_ratio,
                "existing_loan_count": data.existing_loan_count,
                "liquid_assets": data.liquid_assets,
                "liquid_assets_ratio": data.liquid_assets_ratio,
                "bank_relationship_years": data.bank_relationship_years,
                "nsf_count": data.nsf_count,
                "monthly_cash_flow": data.monthly_cash_flow,
                "docs_verified_pct": data.docs_verified_pct,
                "identity_verified": data.identity_verified
            },
            "evaluations": [
                {
                    "factor_id": e.factor_id,
                    "factor_name": e.factor_name,
                    "category": e.category,
                    "value": e.value,
                    "status": e.status.value,
                    "weight": e.weight,
                    "criteria": e.criteria,
                    "notes": e.notes,
                    "risk_score": e.risk_score,
                    "confidence": e.confidence,
                    "source": e.source,
                    "disagreement": e.disagreement
                }
                for e in results.evaluations
            ],
            "summary": {
                "total_factors": results.total_factors,
                "pass_count": results.pass_count,
                "review_count": results.review_count,
                "fail_count": results.fail_count,
                "weighted_score": results.weighted_score,
                "average_risk_score": results.average_risk_score,
                "average_confidence": results.average_confidence
            },
            "decision": {
                "result": decision.decision,
                "rule_applied": decision.rule_applied,
                "rule_description": decision.rule_description,
                "risk_level": decision.risk_level,
                "recommendation": decision.recommendation,
                "average_risk_score": decision.average_risk_score,
                "average_confidence": decision.average_confidence
            }
        }
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"report_{timestamp}.json"
        filepath = self.output_dir / filename
        
        with open(filepath, 'w') as f:
            json.dump(report, f, indent=2)
        
        return str(filepath)
